from typing import List, Dict, Any
from langchain_core.documents import Document
from ..base_loader import BaseMarkitdownLoader

class XlsxLoader(BaseMarkitdownLoader):
    """Loader for XLSX files."""

    def __init__(self, file_path: str, split_by_page: bool = False):
        """Initialize with file path and split_by_page option."""
        super().__init__(file_path)
        self.split_by_page = split_by_page

    def load(self) -> List[Document]:
        """Load and convert XLSX file to Markdown.
        If split_by_page is True, each sheet is a separate document.
        If split_by_page is False, all sheets are combined into a single document.
        """
        try:
            from markitdown import MarkItDown
            converter = MarkItDown()
            result = converter.convert(self.file_path)
            markdown_content = result.text_content
            
            # Create basic metadata
            metadata: Dict[str, Any] = {
                "source": self.file_path,
                "file_name": self._get_file_name(self.file_path),
                "file_size": self._get_file_size(self.file_path),
                "conversion_success": True,
            }
            
            # Extract additional metadata directly from the XLSX file using openpyxl
            try:
                from openpyxl import load_workbook
                workbook = load_workbook(self.file_path, read_only=True, data_only=True)
                
                # Extract document properties if available
                props = workbook.properties
                if hasattr(props, 'creator') and props.creator:
                    metadata["author"] = props.creator
                if hasattr(props, 'title') and props.title:
                    metadata["title"] = props.title
                if hasattr(props, 'subject') and props.subject:
                    metadata["subject"] = props.subject
                if hasattr(props, 'description') and props.description:
                    metadata["description"] = props.description
                if hasattr(props, 'keywords') and props.keywords:
                    metadata["keywords"] = props.keywords
                if hasattr(props, 'category') and props.category:
                    metadata["category"] = props.category
            except ImportError:
                pass
            
            if self.split_by_page:
                documents = []
                # Split Markdown content by sheet headers
                sheet_contents = markdown_content.split("## ")
                for sheet_content in sheet_contents[1:]:
                    lines = sheet_content.splitlines()
                    sheet_name = lines[0].strip()  # First line is the sheet name
                    table_content = '\n'.join(lines[1:])  # Remaining lines are the table
                    
                    page_metadata = metadata.copy()
                    page_metadata["page_number"] = sheet_name
                    doc = Document(page_content=table_content, metadata=page_metadata)
                    documents.append(doc)
                return documents
            else:
                return [Document(page_content=markdown_content, metadata=metadata)]
        except Exception as e:
            # Handle conversion errors
            metadata = {
                "source": self.file_path,
                "file_name": self._get_file_name(self.file_path),
                "conversion_success": False,
                "error": str(e),
            }
            return [Document(page_content="", metadata=metadata)]
